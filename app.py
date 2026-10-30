import io
import os
import re
import secrets
import sys
import time
import uuid
from datetime import datetime, timedelta, timezone
from functools import wraps

import pymysql
import requests
from dotenv import load_dotenv
from flask import Flask, Response, jsonify, render_template, request, send_file
from werkzeug.middleware.proxy_fix import ProxyFix

load_dotenv()
app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1)  # 部署在雲端時取得客人的真實 IP

N8N_URL = os.environ.get("N8N_WEBHOOK_URL", "")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")
DB_NAME = os.environ.get("DB_NAME") or os.environ.get("MYSQLDATABASE") or "maple_orders"
DB = dict(
    host=os.environ.get("DB_HOST") or os.environ.get("MYSQLHOST") or "localhost",
    port=int(os.environ.get("DB_PORT") or os.environ.get("MYSQLPORT") or 3306),
    user=os.environ.get("DB_USER") or os.environ.get("MYSQLUSER") or "root",
    password=os.environ.get("DB_PASSWORD") or os.environ.get("MYSQLPASSWORD") or "",
    charset="utf8mb4",
    cursorclass=pymysql.cursors.DictCursor,
)
TW = timezone(timedelta(hours=8))
STATUSES = ("new", "preparing", "done", "cancelled")


def photo(path, size=240):
    return f"https://images.unsplash.com/{path}?auto=format&fit=crop&w={size}&h={size}&q=75"


HERO_IMG = "https://images.unsplash.com/photo-1535898331935-2d274aff0fbc?auto=format&fit=crop&w=1600&h=700&q=75"

MENU = [
    {"id": "lurou", "cat": "主食", "name": "滷肉飯", "price": 55, "note": "古早味肉燥，淋在白飯上", "img": photo("photo-1679279726946-a158b8bcaa23")},
    {"id": "jitui", "cat": "主食", "name": "雞腿便當", "price": 110, "note": "附三樣配菜", "img": photo("photo-1786201188628-db52cdadb037")},
    {"id": "paigu", "cat": "主食", "name": "排骨便當", "price": 100, "note": "附三樣配菜", "img": photo("photo-1568717099337-5aaecc090548")},
    {"id": "qingcai", "cat": "配菜與湯", "name": "燙青菜", "price": 40, "note": "當日時蔬", "img": photo("photo-1771573042838-8b3adf17562e")},
    {"id": "ludan", "cat": "配菜與湯", "name": "滷蛋", "price": 15, "note": "每日現滷", "img": photo("photo-1587217518657-fa94faff4e7e")},
    {"id": "gongwan", "cat": "配菜與湯", "name": "貢丸湯", "price": 40, "note": "手打貢丸", "img": photo("photo-1706740006781-3ca0149307e5")},
    {"id": "zhenzhu", "cat": "飲料", "name": "珍珠奶茶", "price": 55, "note": "可選大杯", "img": photo("photo-1734770580735-796a00e42cb2")},
    {"id": "lvcha", "cat": "飲料", "name": "無糖綠茶", "price": 30, "note": "冷飲", "img": photo("photo-1637273484093-3e205aed2c59")},
]
MENU_BY_ID = {m["id"]: m for m in MENU}


# ---------- MySQL ----------
def conn():
    return pymysql.connect(database=DB_NAME, **DB)


def init_db():
    try:
        c = pymysql.connect(**DB)
        with c.cursor() as cur:
            cur.execute(f"CREATE DATABASE IF NOT EXISTS `{DB_NAME}` "
                        "CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci")
        c.close()
    except pymysql.err.OperationalError as e:
        if e.args[0] != 1044:  # 1044 = 沒有建立資料庫的權限，直接用現成的資料庫
            raise
    c = conn()
    with c.cursor() as cur:
        cur.execute("""CREATE TABLE IF NOT EXISTS orders (
            id INT AUTO_INCREMENT PRIMARY KEY,
            order_id VARCHAR(10) NOT NULL UNIQUE,
            customer_name VARCHAR(50) NOT NULL,
            phone VARCHAR(12) NOT NULL,
            fulfillment ENUM('pickup','delivery') NOT NULL,
            payment ENUM('cash','card') NOT NULL DEFAULT 'cash',
            address VARCHAR(200) NOT NULL DEFAULT '',
            note VARCHAR(200) NOT NULL DEFAULT '',
            total INT NOT NULL,
            status ENUM('new','preparing','done','cancelled') NOT NULL DEFAULT 'new',
            n8n_sent TINYINT(1) NOT NULL DEFAULT 0,
            created_at DATETIME NOT NULL) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""")
        cur.execute("SHOW COLUMNS FROM orders LIKE 'payment'")
        if not cur.fetchone():  # 舊版資料表補上付款方式欄位
            cur.execute("ALTER TABLE orders ADD COLUMN payment "
                        "ENUM('cash','card') NOT NULL DEFAULT 'cash' AFTER fulfillment")
        cur.execute("""CREATE TABLE IF NOT EXISTS order_items (
            id INT AUTO_INCREMENT PRIMARY KEY,
            order_id VARCHAR(10) NOT NULL,
            item_id VARCHAR(20) NOT NULL,
            item_name VARCHAR(50) NOT NULL,
            price INT NOT NULL,
            qty INT NOT NULL,
            subtotal INT NOT NULL,
            FOREIGN KEY (order_id) REFERENCES orders(order_id) ON DELETE CASCADE) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4""")
    c.commit()
    c.close()


def save_order(o):
    c = conn()
    try:
        with c.cursor() as cur:
            cur.execute(
                "INSERT INTO orders (order_id, customer_name, phone, fulfillment, payment,"
                " address, note, total, created_at) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                (o["order_id"], o["customer"]["name"], o["customer"]["phone"], o["fulfillment"],
                 o["payment"], o["address"], o["note"], o["total"],
                 o["created_at"][:19].replace("T", " ")))
            cur.executemany(
                "INSERT INTO order_items (order_id, item_id, item_name, price, qty, subtotal)"
                " VALUES (%s,%s,%s,%s,%s,%s)",
                [(o["order_id"], i["id"], i["name"], i["price"], i["qty"], i["subtotal"])
                 for i in o["items"]])
        c.commit()
    except Exception:
        c.rollback()
        raise
    finally:
        c.close()


def notify_n8n(order):
    if not N8N_URL:
        return False
    try:
        r = requests.post(N8N_URL, json=order, timeout=10)
        if not r.ok:
            app.logger.error("n8n returned %s: %s", r.status_code, r.text[:300])
        return r.ok
    except requests.RequestException:
        app.logger.exception("n8n unreachable")
        return False


_hits = {}


def too_many(ip, limit=6, window=600):
    now = time.time()
    hits = [t for t in _hits.get(ip, []) if now - t < window]
    if len(hits) >= limit:
        _hits[ip] = hits
        return True
    hits.append(now)
    _hits[ip] = hits
    return False


# ---------- 客人 ----------
@app.get("/")
def index():
    return render_template("index.html", hero=HERO_IMG)


@app.get("/api/menu")
def get_menu():
    return jsonify(MENU)


@app.post("/api/order")
def create_order():
    if too_many(request.remote_addr):
        return jsonify(error="送出太頻繁，請稍後再試"), 429
    d = request.get_json(silent=True) or {}
    name = str(d.get("name", "")).strip()[:50]
    phone = re.sub(r"[\s-]", "", str(d.get("phone", "")))
    note = str(d.get("note", "")).strip()[:200]
    ful = d.get("fulfillment")
    pay = d.get("payment")
    addr = str(d.get("address", "")).strip()[:200]

    if not name:
        return jsonify(error="請填寫姓名"), 400
    if not re.fullmatch(r"\d{8,12}", phone):
        return jsonify(error="電話格式不正確，請輸入 8–12 位數字"), 400
    if ful not in ("pickup", "delivery"):
        return jsonify(error="請選擇取餐方式"), 400
    if ful == "delivery" and not addr:
        return jsonify(error="外送請填寫地址"), 400
    if pay not in ("cash", "card"):
        return jsonify(error="請選擇付款方式"), 400

    items, total = [], 0
    for row in d.get("items", []):
        m = MENU_BY_ID.get(row.get("id"))
        try:
            q = int(row.get("qty", 0))
        except (TypeError, ValueError):
            q = 0
        if not m or not 1 <= q <= 20:
            continue
        total += m["price"] * q
        items.append({"id": m["id"], "name": m["name"], "price": m["price"],
                      "qty": q, "subtotal": m["price"] * q})
    if not items:
        return jsonify(error="請至少選一樣餐點"), 400

    order = {
        "order_id": "A" + uuid.uuid4().hex[:6].upper(),
        "created_at": datetime.now(TW).isoformat(timespec="seconds"),
        "customer": {"name": name, "phone": phone},
        "fulfillment": ful,
        "payment": pay,
        "address": addr if ful == "delivery" else "",
        "note": note,
        "items": items,
        "total": total,
    }

    try:
        save_order(order)  # 資料庫是訂單的唯一記錄
    except Exception:
        app.logger.exception("save order failed")
        return jsonify(error="訂單儲存失敗，請稍後再試"), 500

    if notify_n8n(order):  # 通知失敗不影響訂單，後台會顯示「通知未送出」
        c = conn()
        with c.cursor() as cur:
            cur.execute("UPDATE orders SET n8n_sent=1 WHERE order_id=%s", (order["order_id"],))
        c.commit()
        c.close()

    return jsonify(order_id=order["order_id"], total=total)


# ---------- 店家後台 ----------
def admin_only(f):
    @wraps(f)
    def wrapper(*a, **k):
        if not ADMIN_PASSWORD:
            return Response("請先在 .env 設定 ADMIN_PASSWORD", 503)
        auth = request.authorization
        ok = auth and secrets.compare_digest((auth.password or "").encode(), ADMIN_PASSWORD.encode())
        if not ok:
            return Response("需要登入", 401, {"WWW-Authenticate": 'Basic realm="admin"'})
        return f(*a, **k)
    return wrapper


@app.get("/admin")
@admin_only
def admin():
    return render_template("admin.html")


STATUS_TEXT = {"new": "新訂單", "preparing": "製作中", "done": "已完成", "cancelled": "已取消"}
PAY_TEXT = {"cash": "現金", "card": "刷卡"}
FUL_TEXT = {"pickup": "到店自取", "delivery": "外送"}


def parse_date(value):
    try:
        return datetime.strptime(value or "", "%Y-%m-%d").date().isoformat()
    except ValueError:
        return datetime.now(TW).date().isoformat()


def fetch_orders(cur, status, day):
    """active = 進行中（新訂單＋製作中，不分日期）；其餘依日期查詢。"""
    if status == "active":
        cur.execute("SELECT * FROM orders WHERE status IN ('new','preparing') ORDER BY id ASC LIMIT 200")
    elif status in STATUSES:
        cur.execute("SELECT * FROM orders WHERE status=%s AND DATE(created_at)=%s "
                    "ORDER BY id DESC LIMIT 500", (status, day))
    else:
        cur.execute("SELECT * FROM orders WHERE DATE(created_at)=%s ORDER BY id ASC", (day,))
    orders = cur.fetchall()
    if orders:
        ids = [o["order_id"] for o in orders]
        cur.execute("SELECT * FROM order_items WHERE order_id IN (%s)"
                    % ",".join(["%s"] * len(ids)), ids)
        by = {}
        for it in cur.fetchall():
            by.setdefault(it["order_id"], []).append(it)
        for o in orders:
            o["items"] = by.get(o["order_id"], [])
    return orders


def day_summary(cur, day):
    """營業額只計算「已完成」的訂單。"""
    cur.execute("""SELECT COUNT(*) AS n, COALESCE(SUM(total),0) AS revenue,
        COALESCE(SUM(CASE WHEN payment='cash' THEN total ELSE 0 END),0) AS cash,
        COALESCE(SUM(CASE WHEN payment='card' THEN total ELSE 0 END),0) AS card
        FROM orders WHERE status='done' AND DATE(created_at)=%s""", (day,))
    r = cur.fetchone()
    cur.execute("SELECT COUNT(*) AS n FROM orders WHERE status='cancelled' AND DATE(created_at)=%s", (day,))
    cancelled = cur.fetchone()["n"]
    cur.execute("SELECT COUNT(*) AS n FROM orders WHERE status IN ('new','preparing')")
    pending = cur.fetchone()["n"]
    cur.execute("""SELECT i.item_name, SUM(i.qty) AS qty, SUM(i.subtotal) AS amount
        FROM order_items i JOIN orders o ON o.order_id=i.order_id
        WHERE o.status='done' AND DATE(o.created_at)=%s
        GROUP BY i.item_id, i.item_name ORDER BY qty DESC""", (day,))
    items = [{"name": x["item_name"], "qty": int(x["qty"]), "amount": int(x["amount"])}
             for x in cur.fetchall()]
    return {"date": day, "done": int(r["n"]), "revenue": int(r["revenue"]),
            "cash": int(r["cash"]), "card": int(r["card"]),
            "cancelled": int(cancelled), "pending": int(pending), "items": items}


def safe(v):
    """客人輸入的文字若以 = + - @ 開頭，Excel 會當成公式執行，先加單引號中和。"""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def build_xlsx(day, orders, s):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "當日摘要"
    for row in [["日期", day], ["已完成單數", s["done"]], ["營業額（已完成）", s["revenue"]],
                ["現金", s["cash"]], ["刷卡", s["card"]], ["已取消單數", s["cancelled"]]]:
        ws.append(row)
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    for r in range(3, 6):
        ws.cell(r, 2).number_format = "#,##0"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 16

    wd = wb.create_sheet("訂單明細")
    wd.append(["單號", "時間", "姓名", "電話", "取餐方式", "地址", "付款方式", "狀態", "餐點", "備註", "金額"])
    for o in orders:
        items = "、".join(f'{i["item_name"]}×{i["qty"]}' for i in o["items"])
        wd.append([safe(v) for v in [
            o["order_id"], o["created_at"].strftime("%Y-%m-%d %H:%M"), o["customer_name"], o["phone"],
            FUL_TEXT.get(o["fulfillment"], o["fulfillment"]), o["address"],
            PAY_TEXT.get(o["payment"], o["payment"]), STATUS_TEXT.get(o["status"], o["status"]),
            items, o["note"], o["total"]]])
    for cell in wd[1]:
        cell.font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJK", (10, 17, 12, 14, 10, 28, 10, 10, 40, 24, 10)):
        wd.column_dimensions[col].width = w
    for r in range(2, wd.max_row + 1):
        wd.cell(r, 11).number_format = "#,##0"

    wi = wb.create_sheet("餐點銷量")
    wi.append(["品項", "數量", "金額"])
    for i in s["items"]:
        wi.append([safe(i["name"]), i["qty"], i["amount"]])
    for cell in wi[1]:
        cell.font = Font(bold=True)
    wi.column_dimensions["A"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@app.get("/api/admin/orders")
@admin_only
def admin_orders():
    status = request.args.get("status", "active")
    day = parse_date(request.args.get("date"))
    c = conn()
    try:
        with c.cursor() as cur:
            orders = fetch_orders(cur, status, day)
    finally:
        c.close()
    for o in orders:
        o["created_at"] = o["created_at"].strftime("%m/%d %H:%M")
    return jsonify(orders)


@app.get("/api/admin/summary")
@admin_only
def admin_summary():
    day = parse_date(request.args.get("date"))
    c = conn()
    try:
        with c.cursor() as cur:
            result = day_summary(cur, day)
    finally:
        c.close()
    return jsonify(result)


@app.get("/admin/export.xlsx")
@admin_only
def admin_export():
    day = parse_date(request.args.get("date"))
    c = conn()
    try:
        with c.cursor() as cur:
            orders = fetch_orders(cur, "all", day)
            summary = day_summary(cur, day)
    finally:
        c.close()
    return send_file(io.BytesIO(build_xlsx(day, orders, summary)), as_attachment=True,
                     download_name=f"orders_{day}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.post("/api/admin/orders/<order_id>/status")
@admin_only
def admin_set_status(order_id):
    status = (request.get_json(silent=True) or {}).get("status")
    if status not in STATUSES:
        return jsonify(error="狀態不正確"), 400
    c = conn()
    try:
        with c.cursor() as cur:
            cur.execute("UPDATE orders SET status=%s WHERE order_id=%s", (status, order_id))
        c.commit()
    finally:
        c.close()
    return jsonify(ok=True)


def start_db(retries=10):
    for _ in range(retries):
        try:
            init_db()
            return
        except pymysql.err.OperationalError as e:
            if e.args[0] in (2003, 2006, 2013):  # 資料庫還沒啟動好，等一下再試
                app.logger.warning("MySQL 尚未就緒，3 秒後重試")
                time.sleep(3)
                continue
            sys.exit(f"無法連線 MySQL：{e}\n請確認帳號密碼正確（本機看 .env，雲端看環境變數）。")
    sys.exit("無法連線 MySQL，請確認資料庫服務已啟動。")


start_db()

if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1", port=int(os.environ.get("PORT", 5000)))
